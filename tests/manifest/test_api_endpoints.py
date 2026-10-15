import pytest
import io
import openpyxl
from starlette.testclient import TestClient
from app import app

from auth import clear_rate_limits

@pytest.fixture
def client():
    clear_rate_limits()
    client = TestClient(app)
    # Log in
    client.post('/login', data={'username': 'admin', 'password': 'manifest2026!'})
    return client

def test_index_route(client):
    rv = client.get('/')
    assert rv.status_code == 200
    assert b"Notion Manifest Batch Sync" in rv.content
    assert b"Receiving Lab" in rv.content
    assert b"Manifest Portal" in rv.content

def test_notion_status_route(client):
    rv = client.get('/api/notion/status')
    assert rv.status_code == 200
    data = rv.json()
    assert data["configured"] is True
    assert data["has_api_key"] is True
    assert data["has_database_id"] is True

def test_export_route_multipart(client):
    cases = [{
        "bdo_tag": "BDO-TEST-E2E",
        "test_type": "LEGAL PATERNITY",
        "special_requests": "NONE",
        "receiving_lab": "EUROFINS",
        "is_multipart": True,
        "parts_merged": ["Part 1", "Part 2"],
        "participants": [
            {"name": f"CHILD {i}", "dob": f"0{i}012015", "role": "CHILD", "ethnicity": "AFRICAN"}
            for i in range(1, 6)
        ]
    }]
    rv = client.post('/api/export', json={"cases": cases})
    assert rv.status_code == 200
    assert "spreadsheetml" in rv.headers.get("content-type", "")
    
    # Verify openpyxl can load the returned Excel bytes
    wb = openpyxl.load_workbook(io.BytesIO(rv.content))
    assert "BDO-TEST-E2E" in wb.sheetnames
    ws = wb["BDO-TEST-E2E"]
    assert ws["B2"].value == "BDO-TEST-E2E"
    assert ws["B3"].value == "LEGAL PATERNITY"

def test_update_status_route_success(client, monkeypatch):
    from unittest.mock import MagicMock
    import app

    mock_update = MagicMock(return_value={"total": 2, "updated": 2, "failed": 0, "failed_cases": []})
    monkeypatch.setattr(app, "update_case_status_in_notion", mock_update)

    rv = client.post('/api/notion/update-status', json={"case_page_ids": ["uuid-1", "uuid-2"]})
    assert rv.status_code == 200
    data = rv.json()
    assert data["success"] is True
    assert data["updated"] == 2
    mock_update.assert_called_once_with(case_page_ids=["uuid-1", "uuid-2"], target_status="Dispatched to Laboratory")

def test_update_status_route_from_cases(client, monkeypatch):
    from unittest.mock import MagicMock
    import app

    mock_update = MagicMock(return_value={"total": 1, "updated": 1, "failed": 0, "failed_cases": []})
    monkeypatch.setattr(app, "update_case_status_in_notion", mock_update)

    rv = client.post('/api/notion/update-status', json={
        "cases": [{"bdo_tag": "TAG-1", "case_page_ids": ["uuid-abc"]}]
    })
    assert rv.status_code == 200
    data = rv.json()
    assert data["success"] is True
    assert data["updated"] == 1
    mock_update.assert_called_once_with(case_page_ids=["uuid-abc"], target_status="Dispatched to Laboratory")

def test_update_status_route_empty(client):
    rv = client.post('/api/notion/update-status', json={})
    assert rv.status_code == 200
    data = rv.json()
    assert data["success"] is True
    assert data["total"] == 0

def test_update_status_unauthenticated():
    unauth_client = TestClient(app)
    rv = unauth_client.post('/api/notion/update-status', json={"case_page_ids": ["xyz"]})
    assert rv.status_code == 401
    data = rv.json()
    assert "Authentication required" in data["error"]

def test_favicon_route(client):
    rv = client.get('/favicon.ico')
    assert rv.status_code == 200
    assert "image" in rv.headers.get("content-type", "")
    assert len(rv.content) > 0
