import pytest
import io
import openpyxl
from app import app

@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as client:
        with client.session_transaction() as sess:
            sess['authenticated'] = True
            sess['user'] = 'admin'
        yield client

def test_index_route(client):
    rv = client.get('/')
    assert rv.status_code == 200
    assert b"Notion Manifest Batch Sync" in rv.data
    assert b"Receiving Lab" in rv.data
    assert b"Sync from Notion" in rv.data

def test_notion_status_route(client):
    rv = client.get('/api/notion/status')
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["configured"] is True
    assert data["has_api_key"] is True
    assert data["has_database_id"] is True

def test_export_route_multipart(client):
    cases = [{
        "bdo_tag": "BDO-TEST-E2E",
        "test_type": "LEGAL PATERNITY",
        "special_requests": "NONE",
        "receiving_lab": "EUROFINS",
        "is_multipart": True,
        "parts_merged": ["Part 1", "Part 2"],
        "participants": [
            {"name": f"CHILD {i}", "dob": f"0{i}012015", "role": "CHILD", "ethnicity": "AFRICAN"}
            for i in range(1, 6)
        ]
    }]
    rv = client.post('/api/export', json={"cases": cases})
    assert rv.status_code == 200
    assert rv.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Verify openpyxl can load the returned Excel bytes
    wb = openpyxl.load_workbook(io.BytesIO(rv.data))
    assert "BDO-TEST-E2E" in wb.sheetnames
    ws = wb["BDO-TEST-E2E"]
    assert ws["B2"].value == "BDO-TEST-E2E"
    assert ws["B3"].value == "LEGAL PATERNITY"

def test_update_status_route_success(client, monkeypatch):
    from unittest.mock import MagicMock
    import app

    mock_update = MagicMock(return_value={"total": 2, "updated": 2, "failed": 0, "failed_cases": []})
    monkeypatch.setattr(app, "update_case_status_in_notion", mock_update)

    rv = client.post('/api/notion/update-status', json={"case_page_ids": ["uuid-1", "uuid-2"]})
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["success"] is True
    assert data["updated"] == 2
    mock_update.assert_called_once_with(case_page_ids=["uuid-1", "uuid-2"], target_status="Dispatched to Laboratory")

def test_update_status_route_from_cases(client, monkeypatch):
    from unittest.mock import MagicMock
    import app

    mock_update = MagicMock(return_value={"total": 1, "updated": 1, "failed": 0, "failed_cases": []})
    monkeypatch.setattr(app, "update_case_status_in_notion", mock_update)

    rv = client.post('/api/notion/update-status', json={
        "cases": [{"bdo_tag": "TAG-1", "case_page_ids": ["uuid-abc"]}]
    })
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["success"] is True
    assert data["updated"] == 1
    mock_update.assert_called_once_with(case_page_ids=["uuid-abc"], target_status="Dispatched to Laboratory")

def test_update_status_route_empty(client):
    rv = client.post('/api/notion/update-status', json={})
    assert rv.status_code == 200
    data = rv.get_json()
    assert data["success"] is True
    assert data["total"] == 0

def test_update_status_unauthenticated():
    app.config['TESTING'] = True
    with app.test_client() as unauth_client:
        rv = unauth_client.post('/api/notion/update-status', json={"case_page_ids": ["xyz"]})
        # API routes return 401 JSON when unauthenticated
        assert rv.status_code == 401
        data = rv.get_json()
        assert "Authentication required" in data["error"]

